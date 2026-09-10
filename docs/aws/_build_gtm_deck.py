#!/usr/bin/env python3
"""Build the business / go-to-market AWS deck (Virtusa-styled, no code references).

A concise, executive-facing narrative: the opportunity, the solution on AWS, the
proof, and the engagement path. Distinct from the technical readout in
Galaxy-Agentic-Governance-AWS.pptx. Embeds the reality-based reference diagram.
"""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).resolve().parent
A = HERE / "_assets"
OUT = HERE / "Galaxy-Agentic-Governance-AWS-GTM.pptx"

NAVY = RGBColor(0x16, 0x16, 0x3C); NAVY2 = RGBColor(0x24, 0x22, 0x52)
GREEN = RGBColor(0x2E, 0xE6, 0x74); VIOLET = RGBColor(0x5B, 0x2B, 0xD6)
INK = RGBColor(0x1B, 0x20, 0x30); SUBINK = RGBColor(0x5B, 0x64, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); LIGHT = RGBColor(0xEC, 0xEC, 0xF6)
RULE = RGBColor(0xD3, 0xD7, 0xE3)
COPY = "© 2026 Virtusa Corporation. All Rights Reserved."

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]; _page = {"n": 0}


def _bg(s, c): s.background.fill.solid(); s.background.fill.fore_color.rgb = c
def _rect(s, l, t, w, h, c, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = c
    (sp.line.fill.background() if line is None else setattr(sp.line.color, "rgb", line))
    sp.shadow.inherit = False; return sp
def _text(s, l, t, w, h, runs, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=4):
    tb = s.shapes.add_textbox(l, t, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        for (txt, size, color, bold, italic) in para:
            r = p.add_run(); r.text = txt; r.font.size = Pt(size)
            r.font.color.rgb = color; r.font.bold = bold; r.font.italic = italic; r.font.name = "Calibri"
    return tb
def _wordmark(s, color, right=True):
    l = Inches(11.3) if right else Inches(0.5)
    _text(s, l, Inches(6.95), Inches(1.6), Inches(0.4), [[("virtusa", 16, color, True, False)]],
          align=PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT)
def _footer(s):
    _page["n"] += 1
    _rect(s, Inches(0.5), Inches(6.92), Inches(9.5), Pt(0.75), RULE)
    _text(s, Inches(0.5), Inches(6.98), Inches(9.0), Inches(0.4), [[(f"{_page['n']}    {COPY}", 8.5, SUBINK, False, False)]])
    _wordmark(s, VIOLET, right=True)
def _heading(s, title):
    _text(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.8), [[(title, 26, INK, True, False)]])
    _rect(s, Inches(0.62), Inches(1.15), Inches(1.1), Pt(2.5), GREEN)


def title_slide(title, subtitle, tagline):
    s = prs.slides.add_slide(BLANK); _bg(s, NAVY)
    _rect(s, Inches(11.6), 0, Inches(0.45), SH, NAVY2)
    _rect(s, Inches(12.25), 0, Inches(0.30), SH, VIOLET)
    _rect(s, Inches(12.75), 0, Inches(0.12), SH, GREEN)
    _text(s, Inches(0.7), Inches(0.55), Inches(4), Inches(0.5), [[("virtusa", 20, GREEN, True, False)]])
    _text(s, Inches(0.7), Inches(2.5), Inches(10.4), Inches(2.2), [[(title, 44, WHITE, False, False)]])
    _text(s, Inches(0.72), Inches(4.6), Inches(10.2), Inches(0.9), [[(subtitle, 22, GREEN, True, False)]])
    _text(s, Inches(0.72), Inches(6.4), Inches(9), Inches(0.4), [[(tagline, 14, LIGHT, False, False)]])
    return s


def divider_slide(num, title, blurb=None):
    s = prs.slides.add_slide(BLANK); _bg(s, NAVY)
    _rect(s, 0, Inches(3.0), Inches(0.18), Inches(1.6), GREEN)
    _text(s, Inches(0.7), Inches(2.7), Inches(11.5), Inches(1.2),
          [[(f"{num} · ", 40, GREEN, True, False), (title, 40, WHITE, False, False)]], anchor=MSO_ANCHOR.MIDDLE)
    if blurb:
        _text(s, Inches(0.75), Inches(4.2), Inches(10.8), Inches(1.2), [[(blurb, 16, LIGHT, False, False)]])
    _wordmark(s, GREEN, right=True); return s


def content_slide(title, bullets, *, lead=None, numbered=False):
    s = prs.slides.add_slide(BLANK); _bg(s, WHITE); _heading(s, title)
    paras = []
    if lead:
        paras.append([(lead, 15, SUBINK, False, True)])
    for i, b in enumerate(bullets):
        prefix = f"{i+1}.  " if numbered else "•  "
        if " — " in b:
            head, rest = b.split(" — ", 1)
            paras.append([(prefix + head + " — ", 14.5, INK, True, False), (rest, 14.5, INK, False, False)])
        else:
            paras.append([(prefix + b, 14.5, INK, False, False)])
    _text(s, Inches(0.7), Inches(1.45), Inches(12), Inches(5.2), paras, space=9)
    _footer(s); return s


def table_slide(title, headers, rows, *, lead=None, col_w=None, fsize=11):
    s = prs.slides.add_slide(BLANK); _bg(s, WHITE); _heading(s, title)
    top = Inches(1.45)
    if lead:
        _text(s, Inches(0.7), Inches(1.28), Inches(12), Inches(0.4), [[(lead, 13, SUBINK, False, True)]]); top = Inches(1.9)
    t = s.shapes.add_table(len(rows) + 1, len(headers), Inches(0.6), top, Inches(12.1), Inches(0.3)).table
    if col_w:
        for i, w in enumerate(col_w): t.columns[i].width = Inches(w)
    def cell(c, txt, size, color, bold, fill):
        c.fill.solid(); c.fill.fore_color.rgb = fill; c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame; tf.clear(); tf.word_wrap = True
        r = tf.paragraphs[0].add_run(); r.text = str(txt)
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = "Calibri"
    for j, h in enumerate(headers): cell(t.cell(0, j), h, fsize + 0.5, WHITE, True, NAVY)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell(t.cell(i + 1, j), val, fsize, INK, j == 0, WHITE if i % 2 == 0 else LIGHT)
    _footer(s); return s


def image_slide(title, img, caption=None, *, lead=None, max_w=12.4, max_h=4.7):
    s = prs.slides.add_slide(BLANK); _bg(s, WHITE); _heading(s, title)
    top = 1.5
    if lead:
        _text(s, Inches(0.7), Inches(1.25), Inches(12), Inches(0.5), [[(lead, 13, SUBINK, False, False)]]); top = 1.95
    iw, ih = Image.open(A / img).size; aspect = iw / ih; w = max_w
    if w / aspect > max_h: w = max_h * aspect
    s.shapes.add_picture(str(A / img), Inches((13.333 - w) / 2), Inches(top), width=Inches(w))
    if caption:
        _text(s, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.35), [[(caption, 9.5, SUBINK, False, True)]], align=PP_ALIGN.CENTER)
    _footer(s); return s


def closing_slide():
    s = prs.slides.add_slide(BLANK); _bg(s, NAVY)
    _rect(s, Inches(12.75), 0, Inches(0.12), SH, GREEN)
    _text(s, Inches(0.7), Inches(2.9), Inches(11), Inches(1.5), [[("Let's govern your agents on AWS.", 40, WHITE, False, False)]], anchor=MSO_ANCHOR.MIDDLE)
    _text(s, Inches(0.72), Inches(4.4), Inches(11), Inches(0.5), [[("Galaxy Agentic Governance & Security · Amazon Web Services", 16, GREEN, True, False)]])
    _wordmark(s, GREEN, right=True); return s


# ============================ BUILD ============================
title_slide("Galaxy Agentic Governance & Security",
            "Move enterprise AI agents from pilot to production — safely, on AWS",
            "Solution overview · Amazon Web Services")

content_slide("Executive summary", lead="Galaxy is a runtime governance and security layer for enterprise AI agents, delivered on AWS.", bullets=[
    "The problem it removes — AI agent projects stall before production because their actions are neither governed nor auditable",
    "What it adds — identity, guardrails, data protection, and a tamper-evident audit trail around every agent action",
    "Where it runs — natively on Amazon Bedrock and Bedrock AgentCore, operating on AWS today",
    "Who it is for — enterprises in regulated, high-consequence industries that need agents in production, not pilots",
])

divider_slide("1", "The opportunity", "Why agent governance is a production blocker — and a board-level priority in 2026.")

content_slide("Why agents stall before production", lead="Enterprises are deploying autonomous AI agents, but most cannot leave the pilot stage.", bullets=[
    "Agent actions are real and often irreversible — a wrong step can move money, expose data, or trigger downstream systems",
    "Agent behavior is a black box — leaders cannot see, prove, or control what an agent actually did",
    "Native cloud controls stop at the model — they do not govern the agent's tools, its data access, or agent-to-agent calls",
    "The result — high-value projects stay locked in sandboxes and the value is never realized",
])

content_slide("Why now", lead="Two forces make agent governance urgent this year.", bullets=[
    "Regulation is enforceable — EU AI Act high-risk obligations and DORA now carry material financial penalties",
    "Adoption is outpacing control — most organizations are already running agents, and sensitive data is flowing into them",
    "Autonomy raises the stakes — zero-human-in-the-loop designs need controls that operate without a person in the loop",
    "The window — governance decides which pilots reach production in the current planning cycle",
])

divider_slide("2", "The solution on AWS", "A control plane that makes an autonomous agent safe to run in production, with proof.")

content_slide("What Galaxy delivers", lead="Galaxy makes an autonomous agent safe to run in production — and produces the evidence to prove it.", bullets=[
    "Unblocks production — a control plane that lets agents leave the sandbox with confidence",
    "Identity for every agent — each agent runs as its own least-privilege identity, so every action is attributable",
    "Guardrails at every step — checks before and after the model, and before and after every tool the agent uses",
    "Data protection at the source — sensitive fields are masked or withheld before an agent ever sees them",
    "Tamper-evident audit — a sealed, verifiable record of every decision, ready for auditors and regulators",
    "Governed collaboration — agent-to-agent hand-offs are authorized and logged, never implicit",
])

image_slide("Reference architecture on AWS", "reference-architecture-aws-gtm.png",
            lead="Galaxy owns the security and control plane and maps cleanly onto the AWS agent stack.",
            caption="Green = live on AWS today · amber = reference architecture / roadmap · grey = AWS-native option not required.")

content_slide("Where Galaxy adds value", lead="Galaxy is not another place to build agents — it is the governance layer around them.", bullets=[
    "Security & control plane — the gateway, guardrails, identity, and policy that sit in front of every agent",
    "Governance — compliance mapping, policy lifecycle, and human escalation for exceptions",
    "Enforcement chokepoints — the points where a bad action is stopped, not merely recorded",
    "Audit & telemetry — the evidence trail and the live operational view",
    "Everything else — the agent runtime, model, and infrastructure — stays native AWS; Galaxy supplies the governed agents and the controls on top",
])

table_slide("Two ways to run on AWS", ["Option", "How it works", "Status"], [
    ["Managed (Bedrock AgentCore)", "Agents run on the AWS-managed agent runtime; policy and content controls are enforced at a managed gateway.", "Live on AWS"],
    ["Gateway-proxied (Amazon Bedrock)", "Agents reach the model only through a governed gateway that re-checks every request against policy.", "Reference deployment"],
], lead="Both options run the full control set; the choice is an operating-model decision.", col_w=[3.4, 6.9, 1.8])

content_slide("Native to AWS", lead="Galaxy is built on AWS services, not around them — which makes it straightforward to co-sell and to operate.", bullets=[
    "Model — Amazon Bedrock (Claude family)",
    "Agent runtime & policy — Amazon Bedrock AgentCore: managed runtime, gateway, policy engine, and identity",
    "Identity & secrets — AWS Identity and Access Management, Security Token Service, and Secrets Manager",
    "Audit & telemetry — Amazon DynamoDB, Amazon CloudWatch, and AWS X-Ray",
])

content_slide("Governance and compliance coverage", lead="Controls are mapped to the standards enterprise buyers are measured against.", bullets=[
    "EU AI Act — classification, human oversight, and the audit evidence high-risk systems require",
    "ISO/IEC 42001 and NIST AI RMF — a recognized operating model for AI risk management",
    "OWASP Agentic Top 10 — coverage of the emerging agent-specific threat model",
    "DORA and sector rules — operational resilience and immutable audit for financial and healthcare use",
    "Note — Galaxy provides the controls and evidence that support conformance; it is not itself a certification",
])

content_slide("Proof and status", lead="This is operating on AWS today, not a slideware concept.", bullets=[
    "Live on AWS — the managed AgentCore deployment is operational in a production AWS region",
    "Validated controls — the full control set is exercised on both its success and its denial path in a repeatable matrix",
    "Evidence on demand — every run produces a self-describing conformance report suitable for review",
    "Portable core — the same governance extends to Azure and Google Cloud without re-engineering the agents",
])

content_slide("Where it pays off", lead="Galaxy targets high-consequence, regulated environments where an agent's mistake is expensive.", bullets=[
    "Financial services — autonomous operations with immutable audit and least-privilege data access",
    "Healthcare — sensitive-data protection and human oversight built into the agent, not bolted on",
    "Engineering, procurement & construction — non-human identity for workflows with no person in the loop",
    "Commercial real estate — transaction tracing and activity metrics across automated workflows",
])

divider_slide("3", "Engagement", "A short, structured path from conversation to a governed pilot on AWS.")

content_slide("How we engage", lead="A staged path that reaches a production-ready governed pilot quickly.", bullets=[
    "Architecture workshop — map the client's agent estate onto the AWS reference architecture",
    "Governed pilot — stand up one high-value use case with the full control set on AWS",
    "Production readiness — compliance mapping, audit evidence, and an operating model for scale",
    "AWS co-sell — align to the client's AWS account team and native services throughout",
])

closing_slide()

prs.save(str(OUT))
print(f"wrote {OUT} — {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
