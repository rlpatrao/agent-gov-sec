#!/usr/bin/env python3
"""Build the Virtusa-styled messaging brief .docx with project-accurate answers."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

PURPLE = RGBColor(0x5B, 0x2D, 0x8E)   # virtusa wordmark / sub-headings
MAGENTA = RGBColor(0xD6, 0x00, 0x6D)  # WHY/WHO/WHAT/HOW labels
LEAD = RGBColor(0x1F, 0x3D, 0x7A)     # bold lead-ins
BODY = RGBColor(0x2E, 0x5C, 0x9E)     # blue body text
DARK = RGBColor(0x2B, 0x2B, 0x2B)     # field names
GREY = "C9CEDD"

doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.7); s.bottom_margin = Inches(0.7)
    s.left_margin = Inches(0.8); s.right_margin = Inches(0.8)
n = doc.styles['Normal']; n.font.name = 'Calibri'; n.font.size = Pt(10.5)

def runs(p, lead, text, lead_c=LEAD, body_c=BODY, sz=10.5):
    if lead:
        r = p.add_run(lead + ' '); r.bold = True; r.font.color.rgb = lead_c; r.font.size = Pt(sz)
    if text:
        r = p.add_run(text); r.font.color.rgb = body_c; r.font.size = Pt(sz)

def render(cell, blocks):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    first = True
    for blk in blocks:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        kind = blk[0]
        pf = p.paragraph_format
        if kind == 'h':
            r = p.add_run(blk[1]); r.bold = True; r.font.color.rgb = PURPLE; r.font.size = Pt(11)
            pf.space_before = Pt(7); pf.space_after = Pt(2)
        elif kind == 'p':
            runs(p, blk[1], blk[2]); pf.space_after = Pt(4)
        else:
            ind = 0.20 if kind == 'b' else 0.44
            glyph = '•  ' if kind == 'b' else '–  '
            pf.left_indent = Inches(ind); pf.first_line_indent = Inches(-0.20)
            gr = p.add_run(glyph); gr.font.color.rgb = BODY; gr.font.size = Pt(10.5)
            runs(p, blk[1], blk[2]); pf.space_after = Pt(2)

def label(cell, prefix, name):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = cell.paragraphs[0]
    if prefix:
        r = p.add_run(prefix); r.bold = True; r.font.color.rgb = MAGENTA; r.font.size = Pt(9); r.add_break()
    r = p.add_run(name); r.bold = True; r.font.color.rgb = DARK; r.font.size = Pt(12)

def borders(tbl):
    pr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}')
        x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '6')
        x.set(qn('w:space'), '0'); x.set(qn('w:color'), GREY)
        b.append(x)
    pr.append(b)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); pr.append(lay)

# ── title block ──
h = doc.add_paragraph(); r = h.add_run('virtusa'); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = PURPLE
t = doc.add_paragraph(); r = t.add_run('Messaging Brief'); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = DARK
sub = doc.add_paragraph(); r = sub.add_run('Galaxy Agentic Governance & Security — solution accelerator · internal sales enablement')
r.font.size = Pt(10.5); r.font.color.rgb = BODY
doc.add_paragraph()

ROWS = [
("", "Overview", [
    ('p', None, "This brief covers the technical narrative and internal-enablement assets for the Galaxy Agentic Governance & Security solution accelerator — a runtime governance and security control plane for autonomous and multi-agent AI systems."),
    ('p', None, "The accelerator is built on a cloud-agnostic governance core with a live AWS reference deployment (Amazon Bedrock AgentCore, us-east-2) and a documented path to Azure and GCP. The immediate need is to equip chief architects, sales, and client partners with a structured technical narrative that unblocks production-stage enterprise pipeline conversations and satisfies the internal AWS go-to-market cadence (the “gear” program)."),
]),
("WHY:", "Objective(s)", [
    ('b', "Equip internal teams for high-value client engagement:", "Provide chief architects, sales staff, and client partners with internal enablement tools — specifically a detailed explanatory product video — to confidently initiate complex security discussions with enterprise buyers."),
    ('b', "Establish an enterprise “foot in the door”:", "Use the centralized technical documentation and the AWS reference demo to unblock production-level pipeline conversations, turning the “black box” of agentic security into a tangible competitive differentiator."),
    ('b', "Build long-term solution awareness and credibility:", "Focus this phase on strategic awareness and demand among enterprise technical architects rather than public lead generation, establishing the narrative foundation for a later virtusa.com solution page."),
    ('b', "Develop a native AWS go-to-market motion:", "Satisfy internal cadence and “gear” program requirements by finalizing a reference architecture mapped to AWS-native services (Bedrock AgentCore, Cedar, IAM/STS, DynamoDB, Secrets Manager, X-Ray/CloudWatch), backed by a running AWS reference deployment."),
]),
("WHO:", "Target audience", [
    ('h', "1. Internal Enablement Segment (immediate audience)"),
    ('b', "Buyer / job function:", "Chief architects, sales staff, client partners."),
    ('b', "Behaviors:", "Navigating internal cadence reviews (the “gear” program, AWS-native service requests) and initiating specialized security discussions with high-value clients."),
    ('b', "Attitude:", "Seeking structured clarity to demystify multi-layered security architectures and convert the “black box” of agentic AI into a competitive differentiator that unblocks pipeline accounts."),
    ('h', "2. Enterprise Buyer Segment (ultimate customer)"),
    ('b', "Industry / company:", "Large organizations in high-consequence verticals, reflected in active pipeline: a global EPC firm, a Fortune 500 commercial real estate firm, a global financial firm, and a large healthcare payer."),
    ('b', "Buyer / job function:", "Enterprise buyers and senior technical architects assessing risk, cloud governance, compliance, and target operating models."),
    ('b', "Demographics:", "Cross-functional global technology organizations anchored in AWS or multi-cloud environments."),
    ('b', "Behaviors:", "Among the 79% of organizations deploying or configuring AI agents; working to control shadow-AI exposure (63% of employees pasting sensitive data into AI tools); designing autonomous workflows with non-human identity where no person is in the loop."),
    ('b', "Attitude:", "Treat security as foundational; hesitant to move agents out of preview because agentic security feels uninstrumented; under compliance pressure from the EU AI Act and DORA; aware that agentic threats multiply — a single hallucination can drive an irreversible action, poisoned memory persists across sessions, an unmonitored node expands blast radius."),
]),
("WHO:", "Problem statement", [
    ('p', None, "Enterprise technology leaders are deploying autonomous AI agents but hitting a production ceiling. They face shadow-AI exposure (63% pasting sensitive data into AI tools), cross-session memory poisoning, and compliance pressure from binding law (EU AI Act, up to €35M / 7% of revenue; DORA). Internal sales and architecture teams lack a concrete framework to explain how to instrument, audit, and secure these systems."),
    ('h', "What they think now vs. what we want them to know"),
    ('b', "Currently:", "Agentic AI needs foundational security, but native cloud features (often limited or in preview) are assumed sufficient — or agentic security is seen as an uninstrumented “black box” too risky beyond sandboxes."),
    ('b', "We want them to know:", "Agentic AI introduces a non-deterministic threat model where a hallucination becomes an irreversible action. Cloud providers supply the model and infrastructure but not the orchestration and safety instrumentation needed to control agentic code. The accelerator adds that layer — a single enforcement pipeline (49 controls across four hooks), an out-of-process policy chokepoint under a separate cloud identity, and a tamper-evident audit trail — so agents run in production with every decision attributable and reversible by policy."),
    ('h', "Key narrative pillars"),
    ('b', "Enterprise buyer:", "“I must move our agents out of preview into production, including zero-human-in-the-loop designs, without sacrificing governance or data security.”"),
    ('b', "Internal enablement:", "“I need a structured technical narrative and reference architecture to satisfy the ‘gear’ cadence and confidently advise high-value clients on agentic security.”"),
    ('b', "Because / so that:", "Security is treated as non-negotiable in production agentic discussions; clients can execute real-world automated actions — non-human identity management, transaction tracing — while keeping immutable, hash-chained audit logs."),
    ('b', "However:", "Native cloud tooling supplies compute and preview features but not the integrated instrumentation to monitor and control agentic runtime; traditional security fails against language-based attacks, persistent memory poisoning, and dynamic agent graphs that multiply blast radius."),
]),
("WHAT:", "Key messages (Benefits)", [
    ('h', "Single most important message"),
    ('p', None, "Cloud providers supply the raw model and compute; they do not supply the orchestration and safety instrumentation needed to control autonomous agentic code. The Galaxy accelerator adds that layer — a cloud-agnostic governance core with a live AWS implementation — turning agentic security from an uninstrumented black box into an instrumented, audited, production-ready control plane."),
    ('b', "Benefit 1 — Unblocks production AI safely:", "Introduces an end-to-end security and control plane that lets autonomous projects leave sandboxes; supports zero-human-in-the-loop designs without sacrificing data privacy or governance; gives sales and architecture teams a structured narrative to resolve pipeline blockers."),
    ('b', "Benefit 2 — Cloud-agnostic core with a live AWS reference:", "Core modules import no cloud SDK and no agent framework; mapped across all five architecture layers; live AWS reference deployment on Bedrock AgentCore, with a documented Azure/GCP roadmap."),
    ('b', "Benefit 3 — Neutralizes non-deterministic threats and regulatory risk:", "Real-time circuit breakers, MCP tool gates, memory-access guards, and field-grained data access control (column masking and row filtering on a path that never traverses the LLM) keep a single hallucination from escalating; non-human identity lifecycle with per-agent roles and continuous behavioral monitoring; alignment to DORA and the EU AI Act."),
]),
("WHAT:", "Support (RTBs)", [
    ('h', "RTB 1 — Validated by active enterprise pipeline (Benefit 1)"),
    ('b', "Global EPC firm:", "Addresses the ask for end-to-end security and a non-human identity deep dive for no-person-in-the-loop agents."),
    ('b', "Fortune 500 commercial real estate:", "Supports AI-driven travel & expense automation with comprehensive logging, activity metrics, and end-to-end transaction tracing."),
    ('b', "Global financial & healthcare payer:", "Delivers a multi-cloud target operating model extending identity and security controls across legacy platforms."),
    ('h', "RTB 2 — Instrumented 5-layer architecture with a live AWS deployment (Benefit 2)"),
    ('b', "Five-layer reference model:", "Agent Application, Agent Services, Security & Control Plane, Runtime & Platform, Infrastructure — with cross-cutting governance and observability bands."),
    ('b', "Live AWS deployment:", "Bedrock AgentCore on us-east-2 — per-persona runtimes → MCP gateway → Cedar authorization (ENFORCE) + content-control interceptor Lambdas → tool execution; hash-chained DynamoDB audit ledger; OTel → X-Ray / CloudWatch tracing; 259 tests passing."),
    ('b', "Native AWS mapping:", "Bedrock (claude-sonnet-4-6) for the model; AgentCore Gateway/Runtime/Identity; Cedar policy engine for authorization; GuardPipeline content controls layered over Bedrock Guardrails; IAM/STS and per-agent roles for non-human identity; Secrets Manager for just-in-time credentials; DynamoDB for the immutable ledger. Two deployment methods: a Bedrock proxy (reference Terraform) and AgentCore (live)."),
    ('h', "RTB 3 — Alignment with binding regulatory frameworks (Benefit 3)"),
    ('b', "EU AI Act readiness:", "Audit trails, classification boundaries, and human-oversight gates to address fines up to €35M or 7% of revenue ahead of the August 2026 high-risk enforcement deadline."),
    ('b', "UK GDPR & DORA:", "Automated bias monitoring, DPIA support, and ICT-resilience controls aligned to UK GDPR (Feb 2026) and DORA."),
    ('b', "Agentic standards:", "Built against OWASP Agentic Top 10, NIST AI RMF 1.0, NIST agent standards, CSA Agentic Trust, and ISO/IEC 42001, wired into the standards-crosswalk and control-verification artifacts."),
]),
("HOW:", "Desired action (CTA)", [
    ('p', None, "Because this phase prioritizes internal enablement over public marketing, the calls to action guide internal teams into high-value client engagement:"),
    ('b', "Access the shared repository:", "Review white papers, cross-cutting diagrams, and the AWS reference demo in the shared folder."),
    ('b', "Use the explanatory product video:", "Master the multi-layer agentic security narrative and turn the technical “black box” into a clear differentiator."),
    ('b', "Initiate client architecture discussions:", "Use the framework and reference recordings as a foot in the door for security and cloud-governance conversations."),
    ('b', "Unblock pipeline accounts:", "Apply the 5-layer narrative to resolve production-level security and regulatory blockers across target verticals."),
]),
("HOW:", "Asset list", [
    ('h', "1. Immediate priority — internal enablement"),
    ('b', "Detailed explanatory product video:", "Technical walkthrough of the deployment methods, the 5-layer narrative, and tool accelerators; the prioritized foot-in-the-door asset."),
    ('b', "Solution demos & screen recordings:", "Live AWS (AgentCore) reference implementation; Azure/GCP shown as roadmap."),
    ('b', "Architecture artifacts:", "Reference-architecture diagram, deployment-topology diagram, and the formal AWS architecture document (.docx) in this repository."),
    ('h', "2. Deferred phase — public-facing"),
    ('b', "Future assets:", "Solution webpage (virtusa.com); conceptual marketing video; threat/architecture infographic; targeted social campaigns."),
]),
("HOW:", "Campaign budget", [
    ('p', None, "Internal — to be completed by the requesting team (budget, production-vs-execution split, cost code)."),
]),
("", "Additional comments", [
    ('b', "Key date:", "EU AI Act high-risk enforcement, August 2026 — primary urgency driver."),
    ('b', "Dependency:", "The live AWS demo runs on a specific account/region (us-east-2); keep account identifiers masked in any external asset."),
    ('b', "Must-haves:", "Virtusa branding; AWS partner co-marketing alignment for the GTM/“gear” track."),
]),
("", "Responsible", [('p', None, "Shaiwalini")]),
("", "Approver", [('p', None, "Rajesh Patrao")]),
("", "Consult", [('p', None, "Rajesh Patrao")]),
("", "Inform", [('p', None, "Rajesh Patrao")]),
("", "Support", [('p', None, "Assigned by the internal agency based on capacity and skillset.")]),
]

tbl = doc.add_table(rows=len(ROWS), cols=2)
tbl.allow_autofit = False
borders(tbl)
for i, (prefix, name, blocks) in enumerate(ROWS):
    lc = tbl.cell(i, 0); rc = tbl.cell(i, 1)
    lc.width = Inches(1.55); rc.width = Inches(5.35)
    label(lc, prefix, name)
    render(rc, blocks)

out = "docs/aws/Galaxy-Agentic-Governance-Messaging-Brief.docx"
doc.save(out)
print("wrote", out)
